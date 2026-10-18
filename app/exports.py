"""Exports: an Excel data pack structured like the Commission's Communication
Template sections, and a PDF evidence pack for the verifier."""
from __future__ import annotations

import html
import io
import logging
import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.fonts import addMapping
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table,
                                TableStyle)

GREEN = "1F6F54"
DISCLAIMER = ("Prepared with CarbonTrail from the operator's documents. Actual values must be verified by an "
              "accredited verifier before use in a CBAM declaration. Default values from IR (EU) 2025/2621 as corrected by "
              "IR (EU) 2026/1740; method per IR (EU) 2025/2547. This file mirrors the sections of the Commission's "
              "Communication Template but is not the official template.")


# Hindi and Gujarati on the documents are quoted as written. The PDF embeds Noto Sans (Latin) and Noto Sans
# Devanagari / Gujarati, shaped with HarfBuzz so conjuncts and vowel signs join correctly.
_FONTS = Path(__file__).resolve().parent / "fonts"
_SCRIPTS = {"Deva": ("Devanagari", r"\u0900-\u097F\uA8E0-\uA8FF"), "Gujr": ("Gujarati", r"\u0A80-\u0AFF")}
_SCRIPT_FACE: dict[str, str] = {}
BASE, BASE_BOLD = "Helvetica", "Helvetica-Bold"


def _family(face: str, regular: str, bold: str) -> None:
    pdfmetrics.registerFont(TTFont(face, str(_FONTS / regular), shapable=True))
    pdfmetrics.registerFont(TTFont(face + "-Bold", str(_FONTS / bold), shapable=True))
    for b, f in ((0, face), (1, face + "-Bold")):
        addMapping(face, b, 0, f)
        addMapping(face, b, 1, f)


def _register_fonts() -> None:
    global BASE, BASE_BOLD
    try:
        _family("NotoSans", "NotoSans-Regular.ttf", "NotoSans-Bold.ttf")
        BASE, BASE_BOLD = "NotoSans", "NotoSans-Bold"
        for code, (name, _) in _SCRIPTS.items():
            _family(f"Noto{code}", f"NotoSans{name}-Regular.ttf", f"NotoSans{name}-Bold.ttf")
            _SCRIPT_FACE[code] = f"Noto{code}"
    except Exception:  # noqa: BLE001  the PDF still builds in Helvetica; Indic text would show as boxes
        logging.exception("could not load the PDF fonts")


_register_fonts()
# every word that contains the script goes in that script's font (it also covers digits and punctuation), so
# HarfBuzz shapes the whole word with one font
_WORD_RE = {code: re.compile(rf"[^\s<>]*[{rng}][^\s<>]*") for code, (_, rng) in _SCRIPTS.items()}


# symbols the embedded fonts do not have
_NO_GLYPH = {"→": "->", "≈": "~", "≥": ">=", "≤": "<=", "✓": "yes"}


def _px(text) -> str:
    """Plain text for a reportlab Paragraph: XML-escaped, with Hindi/Gujarati words set in their own font."""
    t = "" if text is None else str(text)
    for a, b in _NO_GLYPH.items():
        t = t.replace(a, b)
    t = html.escape(t, quote=False)
    for code, rx in _WORD_RE.items():
        face = _SCRIPT_FACE.get(code)
        if face:
            t = rx.sub(lambda m: f"<font face='{face}'>{m.group(0)}</font>", t)
    return t


def _qty(v) -> str:
    """12000.0 -> '12,000'; 72.4 -> '72.4'."""
    if v is None:
        return ""
    v = float(v)
    return f"{v:,.0f}" if v == int(v) else f"{v:,.3f}".rstrip("0").rstrip(".")


_RANK = {"high": 0, "medium": 1, "low": 2, "info": 3}


class Markup(str):
    """A table cell that is already reportlab markup (not escaped again)."""


def _merged_findings(run) -> list[dict]:
    """Agent findings plus the automated checks they do not already explain, most severe first."""
    agent = list(run.ctx.agent_findings) if run.ctx else []
    covered = {c for f in agent for c in f.get("covers", [])}
    rest = [f for f in run.findings if f.get("code") not in covered]
    return sorted(agent + rest, key=lambda f: _RANK.get(f.get("severity"), 2))


def _sheet(wb, title, headers, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=GREEN)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append(r)
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, max(12, min(48, len(str(h)) + 6)))
    ws.freeze_panes = "A2"
    return ws


def data_pack_xlsx(run) -> bytes:
    r, p = run.result, run.profile
    ev = run.normalized.evidence
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    rows = [
        ("CarbonTrail CBAM data pack", ""),
        ("Generated", datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC")),
        ("Operator", p.get("company")), ("Installation", p.get("installation") or "Not stated on the documents"),
        ("Country", p.get("country")),
        ("Reporting period", f"{p.get('period_start')} to {p.get('period_end')}"),
        ("Aggregated goods category", ", ".join({g['aggregated_category'] for g in r['goods']})),
        ("Emissions counted", "Direct only (Annex II)" if r["direct_only"] else "Direct and indirect"),
        ("Activity level (t)", round(r["activity_level_t"], 3)),
        ("Attributed direct emissions (tCO2)", round(r["attributed_direct_t"], 3)),
        ("Precursor embedded emissions (tCO2e)", round(r["precursor_emissions_t"], 3)),
        ("Specific embedded emissions SEE (tCO2e/t)", round(r["see"], 4)),
        ("Certificate price used (EUR/t)", f"{r['price']['eur_per_t']} ({r['price']['quarter']})"),
        ("Gross CBAM cost, actual data (EUR)", round(r["totals"]["gross_cost_actual_eur"], 2)),
        ("Gross CBAM cost, EU default values (EUR)", round(r["totals"]["gross_cost_default_eur"] or 0, 2)),
        ("Estimated payable after free allocation, actual data (EUR)",
         "n/a" if r["totals"].get("net_cost_actual_eur") is None else round(r["totals"]["net_cost_actual_eur"], 2)),
        ("Estimated payable after free allocation, EU default values (EUR)",
         "n/a" if r["totals"].get("net_cost_default_eur") is None else round(r["totals"]["net_cost_default_eur"], 2)),
        ("", ""), ("Note", DISCLAIMER),
    ]
    for a, b in rows:
        ws.append([a, b])
        if isinstance(b, float):
            ws.cell(row=ws.max_row, column=2).number_format = "#,##0.00" if "EUR" in a else "#,##0.000#"
    ws["A1"].font = Font(bold=True, size=14, color=GREEN)
    ws.column_dimensions["A"].width = 60
    ws.column_dimensions["B"].width = 90
    for c in ws["B"]:
        c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")

    _sheet(wb, "A_Installation", ["Field", "Value"], [
        ["Operator name", p.get("company")], ["Installation name", p.get("installation") or "Not stated on the documents"],
        ["Address / city", p.get("city") or "Not stated on the documents"], ["Country", p.get("country")],
        ["Production process", "Cold forging, threading, heat treatment of fasteners (one process)"],
        ["Reporting period start", p.get("period_start")], ["Reporting period end", p.get("period_end")],
    ], {1: 32, 2: 60})
    _sheet(wb, "B_EmInst_SourceStreams",
           ["Source stream", "Fuel", "Activity data", "Unit", "Mass (t)", "Energy (GJ)", "EF (tCO2/TJ)", "OxF", "Emissions (tCO2)", "Attributed to process", "Reason", "Evidence"],
           [[f["id"], f["label"], f["quantity"], f["unit"], None if f["mass_t"] is None else round(f["mass_t"], 3), round(f["energy_gj"], 2),
             f["ef_t_per_tj"], 1.0, round(f["co2_t"], 3), "yes" if f["counted"] else "no", f["reason"], ", ".join(f["evidence"])] for f in r["fuels"]],
           {11: 60, 12: 16})
    _sheet(wb, "C_Precursors",
           ["Precursor", "CN", "Supplier", "Consumed (t)", "SEE used (tCO2e/t)", "Basis", "EU default (tCO2e/t)", "Embedded (tCO2e)", "Evidence"],
           [[x["description"], x["cn"], x["supplier"], round(x["consumed_t"], 3), x["see"], x["basis"], x["default_see"], round(x["ee_t"], 3),
             ", ".join(x["evidence"] + x["see_evidence"])] for x in r["precursors"]])
    _sheet(wb, "D_Processes_Goods",
           ["CN", "Description", "Aggregated category", "Production (t)", "EU exports (t)", "SEE (tCO2e/t)", "EU default incl. mark-up", "Gross CBAM cost, actual data (EUR)", "Gross CBAM cost, EU defaults (EUR)"],
           [[g["cn"], g["description"], g["aggregated_category"], round(g["production_t"], 3), round(g["export_eu_t"], 3), round(g["see"], 4),
             None if g["default_marked"] is None else round(g["default_marked"], 4), round(g["gross_cost_actual_eur"], 2),
             None if g["gross_cost_default_eur"] is None else round(g["gross_cost_default_eur"], 2)] for g in r["goods"]])
    _sheet(wb, "Calculation_steps", ["Step", "Formula", "Working", "Result", "Unit", "Rule", "Evidence"],
           [[s["title"], s["formula"], s["working"], round(s["result"], 4), s["unit"], s.get("rule", ""), ", ".join(s.get("evidence", []))] for s in r["steps"]],
           {1: 40, 2: 34, 3: 80, 6: 50})
    _sheet(wb, "Evidence", ["ID", "Document", "Page", "Kind", "Description", "Quantity", "Unit", "Quote on document", "Read confidence"],
           [[e.id, e.doc_name, e.page, e.kind, e.description, e.quantity, e.unit, e.quote, round(e.confidence, 2)] for e in ev.values()],
           {2: 30, 5: 36, 8: 70})
    _sheet(wb, "Findings", ["Severity", "Title", "Detail", "Action"],
           [[f["severity"], f["title"], f["detail"], f.get("action", "")] for f in _merged_findings(run)],
           {2: 40, 3: 90, 4: 50})
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def evidence_pdf(run) -> bytes:
    r, p = run.result, run.profile
    ev = run.normalized.evidence
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm, topMargin=16 * mm, bottomMargin=16 * mm,
                            title="CarbonTrail evidence pack", author="CarbonTrail")
    ss = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=ss["Title"], textColor=colors.HexColor("#" + GREEN), alignment=0, fontSize=20, fontName=BASE_BOLD)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], textColor=colors.HexColor("#" + GREEN), fontName=BASE_BOLD)
    body = ParagraphStyle("b", parent=ss["BodyText"], fontSize=9, leading=12, fontName=BASE, shaping=1)
    small = ParagraphStyle("s", parent=body, fontSize=7.5, leading=9.5, textColor=colors.HexColor("#444444"))

    def tbl(data, widths, header=True):
        t = Table([[Paragraph(c if isinstance(c, Markup) else _px(c), small) for c in row] for row in data], colWidths=widths, repeatRows=1 if header else 0)
        st = [("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#cccccc")), ("VALIGN", (0, 0), (-1, -1), "TOP")]
        if header:
            st += [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8F2EE"))]
        t.setStyle(TableStyle(st))
        return t

    where = " · ".join(_px(x) for x in (p.get("installation"), p.get("city"), p.get("country")) if x)
    S = [Paragraph("CBAM evidence pack", h1),
         Paragraph(f"<b>{_px(p.get('company'))}</b> · {where} · period {p.get('period_start')} to {p.get('period_end')}", body),
         Spacer(1, 6),
         Paragraph(DISCLAIMER, small), Spacer(1, 10),
         tbl([["Specific embedded emissions", "Activity level", "Direct (attributed)", "Precursors", "Excluded"],
              [Markup(f"<b>{r['see']:.4f} tCO2e/t</b>"), f"{r['activity_level_t']:,.3f} t", f"{r['attributed_direct_t']:,.3f} tCO2",
               f"{r['precursor_emissions_t']:,.3f} tCO2e", f"{r['excluded_t']:,.3f} tCO2"]], [36 * mm] * 5),
         Spacer(1, 10), Paragraph("Goods and comparison with EU default values", h2),
         tbl([["CN", "Description", "Produced (t)", "EU exports (t)", "SEE", "EU default incl. mark-up"]] +
             [[g["cn"], g["description"], f"{g['production_t']:,.3f}", f"{g['export_eu_t']:,.3f}", f"{g['see']:.4f}",
               "n/a" if g["default_marked"] is None else f"{g['default_marked']:.4f}"] for g in r["goods"]],
             [18 * mm, 50 * mm, 24 * mm, 24 * mm, 22 * mm, 40 * mm]),
         Spacer(1, 10), Paragraph("Calculation, step by step", h2)]
    for s in r["steps"]:
        S.append(KeepTogether([Paragraph(f"<b>{_px(s['title'])}</b> — <font face='Courier'>{_px(s['formula'])}</font>", body),
                               Paragraph(_px(s["working"]), small),
                               Paragraph(_px(f"Rule: {s.get('rule', '').rstrip('.')}. Evidence: {', '.join(s.get('evidence', [])) or '—'}"
                                             + (f". Note: {s['note']}" if s.get("note") else "")), small), Spacer(1, 5)]))
    S += [PageBreak(), Paragraph("Evidence register", h2),
          Paragraph("Every quantity used above is linked to the exact words on the source document.", small), Spacer(1, 4),
          tbl([["ID", "Document · page", "Value", "Quote on document", "Read conf."]] +
              [[e.id, f"{e.doc_name} · p.{e.page}", f"{_qty(e.quantity)} {e.unit or ''}".strip(), e.quote, f"{e.confidence:.0%}"]
               for e in ev.values()], [12 * mm, 46 * mm, 26 * mm, 76 * mm, 18 * mm]),
          Spacer(1, 10), Paragraph("Findings for review", h2),
          tbl([["Severity", "Finding", "Action"]] + [[f["severity"], Markup(f"<b>{_px(f['title'])}</b>: {_px(f['detail'])}"), f.get("action", "")]
                                                    for f in _merged_findings(run)], [18 * mm, 110 * mm, 50 * mm]),
          Spacer(1, 16),
          tbl([["Verifier review", ""], ["Name / accreditation", ""], ["Date", ""], ["Opinion", ""]], [50 * mm, 128 * mm], header=False)]
    doc.build(S)
    return buf.getvalue()
